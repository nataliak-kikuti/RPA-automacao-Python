from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
import os

nome_arquivo_alunos ="Alunos.xlsx"
planilhaDadosAlunos = load_workbook(nome_arquivo_alunos)

sheet_selecionada = planilhaDadosAlunos["Nomes"]

for linha in range(2, len(sheet_selecionada["A"]) +1):
    arquivoWord = Document("D:\\Usuários\\Natalia\\Desktop\\RPA\\Word\\Certificado2.docx")

    estilo = arquivoWord.styles["Normal"]

    nomeAluno = sheet_selecionada["A%s " % linha].value


    for paragrafo in arquivoWord.paragraphs:
        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    caminhoCertificados = "D:\\Usuários\\Natalia\\Desktop\\RPA\\Word\\Certificados" + nomeAluno + ".docx"

    # Salvando o certificado do aluno
    arquivoWord.save(caminhoCertificados)

print("Certificados gerados com sucesso!")


